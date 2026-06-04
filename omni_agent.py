import os
import json
import requests
import time
import pandas as pd
from urllib.parse import quote  
from datetime import datetime
from openai import OpenAI  # Leveraging standard OpenAI client library for Hugging Face connectivity
from docx import Document
from pptx import Presentation
from pptx.util import Inches
from fpdf import FPDF

# ==========================================================
# ⚙️ HUGGING FACE SERVERLESS CONFIGURATION
# ==========================================================
# Exact repository ID for the dense Gemma 4 31B instruction-tuned variant
HF_MODEL_NAME = "google/gemma-4-31b-it" 
HF_API_URL = "https://api-inference.huggingface.co/v1"

class UniversalAgent:
    def __init__(self, api_key):
        # Point the client directly to the Hugging Face serverless engine
        self.client = OpenAI(
            base_url=HF_API_URL,
            api_key=api_key  # This reads your User Access Token passed from app.py
        )
        self.model = HF_MODEL_NAME
        self.session = requests.Session()
        self.session.headers.update({
            "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36"
        })

    def get_text(self, prompt, is_json=False):
        try:
            messages = []
            if is_json:
                # For structural parsing, we use a system role instruction with the <|think|> trigger 
                # to optimize Gemma 4's chain-of-thought routing for schema compliance.
                messages.append({
                    "role": "system", 
                    "content": "<|think|> You are a strict JSON structural engine. Output raw JSON ONLY. No formatting markdown backticks (such as ```json). Drop all introduction text or conversational pleasantries."
                })
            messages.append({"role": "user", "content": prompt})

            response = self.client.chat.completions.create(
                model=self.model,
                messages=messages,
                response_format={"type": "json_object"} if is_json else None,
                # Gemma 4 works beautifully with standard low temperature bounds when generating schemas
                temperature=0.1 if is_json else 0.7 
            )
            return response.choices[0].message.content
        except Exception as e:
            return "{}" if is_json else f"Error: {e}"

    # --- 🎨 IMAGE GEN (FOR PPT/PDF) ---
    def _generate_temp_image(self, prompt):
        try:
            encoded_prompt = quote(prompt)
            url = f"[https://gen.pollinations.ai/image/](https://gen.pollinations.ai/image/){encoded_prompt}?width=800&nologo=true"
            r = self.session.get(url, timeout=15)
            if r.status_code == 200:
                temp_name = f"temp_{int(time.time())}.jpg"
                with open(temp_name, "wb") as f:
                    f.write(r.content)
                return temp_name
            return None
        except Exception as e: 
            print(f"DEBUG: Image Error - {e}")
            return None

    # --- 📄 PDF GENERATION ---
    def create_pdf(self, topic, filename):
        try:
            if not filename.endswith(".pdf"): filename = f"{topic.replace(' ', '_')}_{int(time.time())}.pdf"
            pdf = FPDF()
            pdf.add_page()
            pdf.set_font("helvetica", 'B', 16)
            pdf.cell(0, 10, txt=topic.upper(), ln=True, align='C')
            
            img = self._generate_temp_image(f"A professional visual representing {topic}")
            if img:
                pdf.image(img, x=10, y=30, w=180)
                pdf.set_y(150)
                os.remove(img)

            pdf.set_font("helvetica", size=12)
            raw_text = self.get_text(f"Write a comprehensive 3-paragraph report on {topic}")
            clean_text = raw_text.encode('latin-1', 'ignore').decode('latin-1')
            pdf.multi_cell(0, 10, txt=clean_text)
            
            pdf.output(filename)
            return {"message": f"✅ PDF Document Created: {topic}", "file_path": filename}
        except Exception as e: return {"message": f"❌ PDF Error: {e}", "file_path": None}

    # --- 📽️ PPT GENERATION ---
    def create_ppt(self, topic, filename, slides=3):
        try:
            if not filename.endswith(".pptx"): filename = f"{topic.replace(' ', '_')}_{int(time.time())}.pptx"
            prs = Presentation()
            for i in range(slides):
                slide = prs.slides.add_slide(prs.slide_layouts[1])
                slide.shapes.title.text = f"{topic} - Slide {i+1}"
                slide.placeholders[1].text = self.get_text(f"Generate educational presentation bullet points for {topic}, subsection breakdown {i+1}")
            prs.save(filename)
            return {"message": f"✅ Presentation Slides Created: {topic}", "file_path": filename}
        except Exception as e: return {"message": f"❌ PPT Error: {e}", "file_path": None}

    # --- 📊 EXCEL GENERATION ---
    def create_excel(self, topic, filename):
        try:
            if not filename.endswith(".xlsx"): filename = f"{topic.replace(' ', '_')}_{int(time.time())}.xlsx"
            raw_text = self.get_text(f"Provide numerical data about '{topic}'. Structure your output to match this specific schema architecture template layout: {{\"cols\":[\"DataHeader1\",\"DataHeader2\"],\"rows\":[[\"Value1\",\"Value2\"]]}}", is_json=True)
            data = json.loads(raw_text)
            df = pd.DataFrame(data['rows'], columns=data['cols'])
            df.to_excel(filename, index=False)
            return {"message": f"✅ Excel Spreadsheet Created: {topic}", "file_path": filename}
        except Exception as e: return {"message": f"❌ Excel Error: {e}", "file_path": None}

    # --- 📝 WORD GENERATION ---
    def create_word(self, topic, filename):
        try:
            if not filename.endswith(".docx"): filename = f"{topic.replace(' ', '_')}_{int(time.time())}.docx"
            doc = Document()
            doc.add_heading(topic, 0)
            doc.add_paragraph(self.get_text(f"Write a detailed summary analysis report focusing on {topic}"))
            doc.save(filename)
            return {"message": f"✅ Word Document Created: {topic}", "file_path": filename}
        except Exception as e: return {"message": f"❌ Word Error: {e}", "file_path": None}

    # --- 🧠 DISPATCHER ---
    def handle_request(self, user_prompt):
        brain_p = f"""
        User Prompt: "{user_prompt}"
        Isolate the objective action parameters. Choose one target mechanism format: pdf, ppt, excel, word, text.
        Return raw JSON containing your output mappings built precisely to match this footprint: {{"tool": "target_mechanism", "subject": "extracted_topic", "file": "clean_filename"}}
        """
        try:
            raw_res = self.get_text(brain_p, is_json=True)
            res = json.loads(raw_res)
            
            t = res.get('tool', 'text')
            s = res.get('subject', user_prompt)
            f = res.get('file', 'output')

            if t == 'pdf': result = self.create_pdf(s, f)
            elif t == 'ppt': result = self.create_ppt(s, f)
            elif t == 'excel': result = self.create_excel(s, f)
            elif t == 'word': result = self.create_word(s, f)
            else: 
                result = {"message": self.get_text(user_prompt), "file_path": None}

            return json.dumps(result)
            
        except (json.JSONDecodeError, Exception):
            chat_reply = self.get_text(user_prompt)
            return json.dumps({"message": chat_reply, "file_path": None})

if __name__ == "__main__":
    HF_TOKEN = "hf_xxxxxxxxxxxxxxxxxxxxxxxxxxxxxxxxx"  # Replace with token for localized manual testing
    agent = UniversalAgent(HF_TOKEN)
    print("--- 🤖 Omni-Agent (Hugging Face Gemma 4 31B IT Dense Backend) ---")
    while True:
        inp = input("\nYou: ").strip()
        if not inp: break
        response_data = json.loads(agent.handle_request(inp))
        print(f"Agent Message: {response_data['message']}")